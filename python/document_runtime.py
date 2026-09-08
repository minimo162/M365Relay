"""Offline document tools. All writes target new files; originals stay unchanged."""
import argparse
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from contextlib import closing


def encode_result(value):
    import datetime
    def serialise(item):
        if isinstance(item, (datetime.datetime, datetime.date, datetime.time)):
            return item.isoformat()
        raise TypeError(f"Unsupported JSON value: {type(item).__name__}")
    data = json.dumps(value, ensure_ascii=True, allow_nan=False, default=serialise)
    if len(data.encode("utf-8")) > 32 * 1024 * 1024:
        raise ValueError("Document output exceeds 32 MiB; select fewer pages")
    return data


def emit(value):
    print(encode_result(value))


def publish(path, writer):
    requested = Path(path).absolute()
    # Resolve directory aliases, but never follow an existing final symlink:
    # a dangling link must not redirect a new output to an unrelated target.
    path = requested.parent.resolve() / requested.name
    if path.exists() or path.is_symlink():
        raise FileExistsError("Output exists; choose a new filename")
    # Same directory/volume; exclusive final creation avoids overwriting races.
    with tempfile.TemporaryDirectory(prefix=".relay-", dir=path.parent) as temp:
        staged = Path(temp) / path.name
        writer(staged)
        os.link(staged, path)


def selected_pages(spec, total):
    if spec is None:
        return list(range(1, total + 1))
    import re
    if not re.fullmatch(r"\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*", spec):
        raise ValueError("Invalid page selection")
    pages = set()
    for part in spec.split(","):
        bounds = [int(n) for n in part.split("-")]
        first, last = bounds[0], bounds[-1]
        if not 1 <= first <= last <= total:
            raise ValueError("Page selection outside document")
        pages.update(range(first, last + 1))
    return sorted(pages)


def pdf_read(path, selection, text_only=False):
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    result = []
    with closing(pdfium.PdfDocument(path)) as pdf:
        total = len(pdf)
        if text_only and selection is None:
            selection = f"1-{min(5, total)}" if total else None
        for number in selected_pages(selection, total):
            with closing(pdf[number - 1]) as page:
                width, height = page.get_size()
                with closing(page.get_textpage()) as textpage:
                    items = []
                    # PDFium character bounds are authoritative geometry. Do
                    # not invent table cells or estimate boxes from text length.
                    for i in range(0 if text_only else textpage.count_chars()):
                        code = raw.FPDFText_GetUnicode(textpage, i)
                        if not code:
                            continue
                        text = chr(code)
                        if text.isspace():
                            continue
                        left, bottom, right, top = textpage.get_charbox(i)
                        items.append(dict(text=text, x=left, y=height-top,
                                          width=right-left, height=top-bottom,
                                          fontSize=raw.FPDFText_GetFontSize(textpage, i)))
                    result.append(dict(pageNum=number, width=width, height=height,
                                       text=textpage.get_text_range(), textItems=items))
    textless = [p["pageNum"] for p in result if not p["text"].strip()]
    if text_only:
        pages = [dict(pageNum=p["pageNum"], text=p["text"]) for p in result]
        return dict(schema="m365-relay-pdf-text-v1", totalPages=total,
                    parsedPageNumbers=[p["pageNum"] for p in pages],
                    documentComplete=len(pages) == total, pages=pages,
                    pagesWithoutText=textless, ocrEnabled=False,
                    notice="Only listed pages were read. Text follows PDF reading order; table layout is not reconstructed. Empty text does not prove a blank page.")
    return dict(schema="m365-relay-pdf-v1", parser="pypdfium2", parserVersion=str(pdfium.PYPDFIUM_INFO),
                ocrEnabled=False, totalPages=total, selectionComplete=True,
                parsedPageNumbers=[p["pageNum"] for p in result], pagesWithoutText=textless,
                pageErrors=[], warnings=([dict(code="no_text_extracted", pages=textless,
                message="Blank or scanned pages: no OCR was performed.")] if textless else []),
                markdown="\n\n".join(p["text"] for p in result), pages=result,
                notice="Text preserves PDFium reading order, not a reconstructed Markdown table. textItems contain character bounds in top-left page coordinates.")


def pdf_render(path, output, number, scale):
    import pypdfium2 as pdfium
    if not 0.25 <= scale <= 4:
        raise ValueError("Scale must be between 0.25 and 4")
    def write(target):
        with closing(pdfium.PdfDocument(path)) as pdf:
            selected_pages(str(number), len(pdf))
            with closing(pdf[number - 1]) as page:
                w, h = page.get_size()
                if w*h*scale*scale > 40_000_000:
                    raise ValueError("Rendered image too large")
                bitmap = page.render(scale=scale, force_bitmap_format=pdfium.raw.FPDFBitmap_BGRA)
                try:
                    # PDFium already renders pixels. PNG needs only stdlib
                    # zlib/CRC encoding; no separate image codec DLLs.
                    import struct
                    import zlib
                    compressor = zlib.compressobj()
                    compressed = []
                    pixels = bytes(bitmap.buffer)
                    for row in range(bitmap.height):
                        bgra = pixels[row*bitmap.stride:row*bitmap.stride+bitmap.width*4]
                        rgba = bytearray(len(bgra))
                        rgba[0::4], rgba[1::4], rgba[2::4], rgba[3::4] = bgra[2::4], bgra[1::4], bgra[0::4], bgra[3::4]
                        compressed.append(compressor.compress(b"\0" + rgba))
                    compressed.append(compressor.flush())
                    def chunk(kind, data):
                        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind+data))
                    target.write_bytes(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", bitmap.width, bitmap.height, 8, 6, 0, 0, 0)) + chunk(b"IDAT", b"".join(compressed)) + chunk(b"IEND", b""))
                finally:
                    bitmap.close()
    publish(output, write)
    return dict(output=str(Path(output).resolve()), page=number)


def validate_xlsx_input(data):
    import math
    import re
    if not isinstance(data, dict) or set(data) - {"sheet", "cells"} or "cells" not in data:
        raise ValueError("Expected sheet and cells fields")
    name = data.get("sheet", "Sheet1")
    if not isinstance(name, str) or not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name):
        raise ValueError("Invalid sheet name")
    if not isinstance(data["cells"], dict) or len(data["cells"]) > 100000:
        raise ValueError("Invalid or oversized cells object")
    for address, entry in data["cells"].items():
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", address):
            raise ValueError("Expected uppercase single-cell addresses")
        column = 0
        for char in re.match(r"[A-Z]+", address).group():
            column = column*26 + ord(char)-64
        if column > 16384 or int(re.search(r"\d+", address).group()) > 1048576:
            raise ValueError("Cell address outside Excel limits")
        if isinstance(entry, dict):
            if set(entry) not in ({"value"}, {"formula"}):
                raise ValueError("Each cell must specify value OR formula, not both")
            value = next(iter(entry.values()))
            if "formula" in entry and (not isinstance(value, str) or not value or value.startswith("==") or not (value[1:] if value.startswith("=") else value).strip()):
                raise ValueError("Invalid formula string")
        else:
            value = entry
        if value is not None and not isinstance(value, (str, int, float, bool)):
            raise ValueError("Unsupported cell value")
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("Non-finite cell number")
        if isinstance(value, str) and len(value.encode("utf-16-le"))//2 > 32767:
            raise ValueError("Cell text exceeds Excel capacity; split it explicitly instead of truncating")
    return data


def xlsx_create(output, source):
    if Path(source).stat().st_size > 32*1024*1024:
        raise ValueError("Input JSON exceeds 32 MiB; use selected ranges")
    def unique_fields(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ValueError("Duplicate JSON field")
            result[key] = value
        return result
    data = validate_xlsx_input(json.loads(Path(source).read_text(encoding="utf-8-sig"), object_pairs_hook=unique_fields))
    import openpyxl
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = data.get("sheet", "Sheet1")
    for address, entry in data["cells"].items():
        if isinstance(entry, dict) and set(entry) == {"formula"}:
            sheet[address] = entry["formula"] if entry["formula"].startswith("=") else "=" + entry["formula"]
        else:
            value = entry["value"] if isinstance(entry, dict) else entry
            sheet[address] = value
            if isinstance(value, str):
                sheet[address].data_type = "s"
    publish(output, lambda target: wb.save(target))
    wb.close()
    return dict(output=str(Path(output).resolve()), formulasCalculated=False)


def xlsx_read(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in wb:
            if sheet.max_row*sheet.max_column > 100000:
                raise ValueError("Sheet too large for full JSON; use a Python script with a selected range")
            cache = cached[sheet.title]
            values = []
            for row in sheet:
                for cell in row:
                    if cell.value is not None:
                        values.append(dict(address=cell.coordinate, value=cell.value,
                            type=cell.data_type, cachedValue=cache[cell.coordinate].value))
            sheets.append(dict(name=sheet.title, cells=values))
        return dict(sheets=sheets, notice="Formula caches can be missing or stale. openpyxl does not calculate formulas.")
    finally:
        wb.close()
        cached.close()


def office_text(path):
    """Read DOCX/PPTX text using the standard library, with ZIP size bounds."""
    if Path(path).suffix.lower() not in (".docx", ".pptx"):
        raise ValueError("office-text accepts docx/pptx only")
    with zipfile.ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 128*1024*1024:
            raise ValueError("Expanded document too large")
        if Path(path).suffix.lower() == ".docx":
            names = ["word/document.xml"]
        else:
            import posixpath
            from urllib.parse import unquote
            pns = "http://schemas.openxmlformats.org/presentationml/2006/main"
            rns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            relns = "http://schemas.openxmlformats.org/package/2006/relationships"
            presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
            relationships = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
            rels = {}
            for rel in relationships.findall(f"{{{relns}}}Relationship"):
                key = rel.get("Id")
                if not key or key in rels:
                    raise ValueError("Missing or duplicate presentation relationship")
                rels[key] = rel
            names = []
            for slide in presentation.findall(f"{{{pns}}}sldIdLst/{{{pns}}}sldId"):
                rel = rels.get(slide.get(f"{{{rns}}}id"))
                if rel is None or rel.get("Type") != rns + "/slide" or rel.get("TargetMode", "Internal") != "Internal":
                    raise ValueError("Invalid slide relationship")
                target = unquote(rel.get("Target", ""))
                if not target or "\\" in target or ":" in target or "?" in target or "#" in target:
                    raise ValueError("Invalid slide part target")
                name = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "ppt/" + target)
                if name.startswith("../") or name not in archive.namelist():
                    raise ValueError("Missing slide part")
                names.append(name)
        if not names:
            raise ValueError("No document/slide parts found")
        parts = []
        for name in names:
            namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main" if name.startswith("word/") else "http://schemas.openxmlformats.org/drawingml/2006/main"
            paragraphs = []
            for paragraph in ET.fromstring(archive.read(name)).iter(f"{{{namespace}}}p"):
                paragraphs.append("".join(e.text or "" if e.tag == f"{{{namespace}}}t" else
                    "\t" if e.tag == f"{{{namespace}}}tab" else "\n" if e.tag == f"{{{namespace}}}br" else ""
                    for e in paragraph.iter()))
            parts.append(dict(part=name, text="\n".join(paragraphs)))
        return dict(parts=parts, notice="Main-body or slide paragraphs only. Headers, notes, images and embedded objects are not extracted.")


def office_native(operation, source, output):
    if os.name != "nt":
        raise RuntimeError("Installed Windows desktop Office is required")
    expected = {"office-pdf": ".pdf", "xlsx-recalculate": ".xlsx", "docx-create": ".docx", "pptx-create": ".pptx"}[operation]
    if Path(output).suffix.lower() != expected:
        raise ValueError(f"Output must use {expected}")
    if operation in ("docx-create", "pptx-create"):
        if Path(source).stat().st_size > 4*1024*1024:
            raise ValueError("Input JSON too large")
        data = json.loads(Path(source).read_text(encoding="utf-8-sig"))
        blocks = [data] if operation == "docx-create" else data["slides"]
        if not isinstance(blocks, list) or not 1 <= len(blocks) <= 1000:
            raise ValueError("Invalid document blocks")
        for block in blocks:
            if not isinstance(block.get("paragraphs"), list) or not all(isinstance(p, str) for p in block["paragraphs"]):
                raise ValueError("paragraphs must be a list of strings")
            if operation == "pptx-create" and not isinstance(block.get("title"), str):
                raise ValueError("Slide title must be a string")
    ps = Path(os.environ["SystemRoot"]) / "System32/WindowsPowerShell/v1.0/powershell.exe"
    helper = Path(__file__).with_name("Office-Native.ps1")
    def write(target):
        state = target.parent / "office-owned.json"
        command = [str(ps), "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                        "-File", str(helper), "-Operation", operation, "-InputPath", str(Path(source).resolve()),
                        "-OutputPath", str(target), "-StatePath", str(state)]
        try:
            subprocess.run(command, check=True, timeout=120,
                           creationflags=subprocess.CREATE_NO_WINDOW, stdout=subprocess.DEVNULL)
        finally:
            subprocess.run([str(ps), "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(helper),
                            "-Operation", "cleanup", "-StatePath", str(state)], check=True, timeout=15,
                           creationflags=subprocess.CREATE_NO_WINDOW, stdout=subprocess.DEVNULL)
    publish(output, write)
    return dict(output=str(Path(output).resolve()), engine="installed-Microsoft-Office")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("pdf-read"); p.add_argument("input"); p.add_argument("output", nargs="?"); p.add_argument("--pages")
    p = sub.add_parser("pdf-text", help="Read text without coordinates; defaults to first 5 pages"); p.add_argument("input"); p.add_argument("--pages")
    p = sub.add_parser("pdf-render"); p.add_argument("input"); p.add_argument("output"); p.add_argument("--page", type=int, default=1); p.add_argument("--scale", type=float, default=1.5)
    p = sub.add_parser("xlsx-create"); p.add_argument("output"); p.add_argument("data")
    for command in ("xlsx-read", "office-text"):
        p = sub.add_parser(command); p.add_argument("input")
    for command in ("office-pdf", "xlsx-recalculate", "docx-create", "pptx-create"):
        p = sub.add_parser(command); p.add_argument("input"); p.add_argument("output")
    a = parser.parse_args()
    if a.command == "pdf-text":
        result = pdf_read(a.input, a.pages, text_only=True)
        if sum(len(p["text"]) for p in result["pages"]) > 48000:
            raise ValueError("Text exceeds 48000 characters; select fewer pages with --pages")
    elif a.command == "pdf-read":
        result = pdf_read(a.input, a.pages)
        if a.output:
            data = encode_result(result)
            publish(a.output, lambda target: target.write_text(data + "\n", encoding="utf-8"))
            result = dict(output=str(Path(a.output).resolve()), totalPages=result["totalPages"],
                          parsedPageNumbers=result["parsedPageNumbers"],
                          pagesWithoutText=result["pagesWithoutText"],
                          selectionComplete=result["selectionComplete"])
    elif a.command == "pdf-render": result = pdf_render(a.input, a.output, a.page, a.scale)
    elif a.command == "xlsx-create": result = xlsx_create(a.output, a.data)
    elif a.command == "xlsx-read": result = xlsx_read(a.input)
    elif a.command == "office-text": result = office_text(a.input)
    else: result = office_native(a.command, a.input, a.output)
    emit(result)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(f"Document operation failed: {type(error).__name__}: {error}", file=sys.stderr)
        sys.exit(1)
